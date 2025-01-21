from django.shortcuts import render, redirect
from django.http import HttpRequest, HttpResponse
from Excel.forms import ExcelForm,VerifyForm
from django.db.models import Q
from Main.tools import *
from Excel.models import *
from django.urls import reverse
from django.conf import settings
from User.models import is_manager, is_uploader
import pandas as pd
from io import BytesIO,StringIO
from typing import Any
import json
import openpyxl as op
from bson.objectid import ObjectId
from tools.get_image import get_image_data
from Excel.templatetags.bad_image import bad_image
from tools.get_image import compress_image
from tools.encrypt import encrypt_data,decrypt_data
from base64 import b64encode, b64decode
from User.models import get_post
import typing
from PIL import Image
import re
from django.utils import timezone
from datetime import datetime

def mongo_setup_failed(req, **kwargs):
    context = {'alert':'Excel Sheet is empty or MongoDB connection failed'}
    context.update(kwargs)
    return render(req,'Excel/index.html',context)


# Create your views here.
def upload_screen(req:HttpRequest):
    if(not ((req.user.is_authenticated) and (is_uploader(req.user)))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    if(req.method=="POST"):
        f = ExcelForm(req.POST,req.FILES)
        
        connection = MongoConnection(settings.MONGO_URL)
        excel = connection.connect(settings.MONGO_CRED)
        pd_data = None
        
        if(excel is None):
            return render(req,'Excel/upload.html',{'form':f,'alert':'MongoDB connection failed'})
        
        if(f.is_valid()):
            excel_file = req.FILES["file"]
            file_name = f.cleaned_data.get("file_name")
        
            with excel_file.open() as file:
                image_idx, pd_data = image_load(file.read())
                rows, _ = pd_data.shape
                pd_data = pd_data.to_json()
            
            if(pd_data is None):
                return mongo_setup_failed(req,form=f)
            
            verify_idx = {str(i):None for i in range(rows)}
            feedback_idx = {str(i):"" for i in range(rows)}
            lock = {str(i):{'status':False,'time':None} for i in range(rows)}
            
            template = MongoTemplate(file_name,json.loads(pd_data),req.user.username,verify_idx,feedback_idx,lock,image_idx).get_json()
            res = excel.insert_one(template)
            
            verify = VerificationTable(mongo_id=res.inserted_id,belongs=req.user.uploader)
            verify.save()
            
            
            return render(req,'Excel/upload.html',{'form':f,'alert':'Excel Extraction Success'})
        
        return render(req,'Excel/upload.html',{'form':f,'alert':'Something Went Wrong'})
        
    else:
        f = ExcelForm(initial={'username':req.user.username})
    
    return render(req,'Excel/upload.html',{'form':f})

def read_screen(req: HttpRequest) -> HttpResponse:
    if(not ((req.user.is_authenticated) and (is_manager(req.user)))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')

    # simulate read card
    try:
        with open(settings.MEDIA_ROOT + '/compress.txt','r') as f:
            data:dict[str,dict[str,str]] = decrypt_data(settings.KEY,f.read())
    except Exception as e:
        return redirect(reverse('Excel:dash') + f'?alert=Failed to read card ({e})')
        
        
    profile_img = r'^Profile_Image$'
    sem_data = r'.+Sem_(\d+)$'
    
    result, header = data.get('data',{}), data.get('header',{})
    
    columns = result.keys()
    
    profile_col = [i for i in columns if re.match(profile_img,i)]
    sem_col = [i for i in columns if re.match(sem_data,i)]
    personal_col = [i for i in columns if((i not in profile_col) and (i not in sem_col))]
    
    sem_dict:dict[str,list[str]] = {}
    
    profile_col = profile_col[0] if profile_col else None
    
    for i in sem_col:
        sem:list[str] = re.findall(sem_data,i)
        
        if(sem):
            sem = sem[0]
            if(sem not in sem_dict): sem_dict[sem] = list()
            sem_dict[sem].append(i)

    view = ReportStructure(profile_img=profile_col,personal_info=personal_col,sem_data=sem_dict)

    return render(req,'Excel/read.html',context={'data':result,'personal':view.personal_info,'pic':view.profile_img,'sem_dict':view.sem_data,**header})


def image_load(data: bytes) -> tuple[list[str], pd.DataFrame]:
    
    pd_data:pd.DataFrame = None
    op_data:op.Workbook = None
    converted:set[str] = set()
    
    try:
        pd_data = pd.read_excel(BytesIO(data))
        op_data = op.load_workbook(BytesIO(data))
    except:
        return converted, pd_data
        

    if(op_data and len(op_data.sheetnames)==0):
        return converted, pd_data
        

    op_sheet = op_data[op_data.sheetnames[0]]
    image = get_image_data(op_sheet)
    
    for i, row in pd_data.iterrows():
        for j, col in enumerate(row):
            pd_data[pd_data.columns[j]] = pd_data[pd_data.columns[j]].astype(str)
            if((i+1,j) in image):
                converted.add(pd_data.columns[j])
                    
                pd_data.iat[i,j] = image[(i+1,j)]
                    
    return list(converted), pd_data
        

def dash_board(req: HttpRequest) -> HttpResponse:
    if(not ((req.user.is_authenticated) and (is_uploader(req.user) or is_manager(req.user)))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    user = req.user
    username = req.user.username
    
    excel = MongoConnection(settings.MONGO_URL).connect(settings.MONGO_CRED)
    
    if(excel is None):
        return mongo_setup_failed(req)
        
    available_id = [ ObjectId(i.mongo_id) for i in VerificationTable.objects.only('mongo_id')]
        
    verify = excel.find({"_id":{"$in":available_id},"verify.user":{"$eq":username}}).to_list()
    upload = excel.find({"_id":{"$in":available_id},"belongs":{"$eq":username}}).to_list()

    available_id = [ ObjectId(i.mongo_id) for i in VerificationTable.objects.only('mongo_id')]
    
    if(is_manager(user)):
        return render(req,'Excel/dash/manager.html',{'verify':verify})
        
    elif(is_uploader(user)):
        return render(req,'Excel/dash/uploader.html',{'upload':upload})



def data_view(req: HttpRequest,id) -> HttpResponse:
    if(not ((req.user.is_authenticated) and (is_uploader(req.user) or is_manager(req.user)))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    result = get_data(id,req.user.username)
    if(result is None):
        return redirect(reverse('Excel:dash')+'?alert=Record not found')

    pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
    image_idx:list = result.get('excel',{}).get('image',[])

    # removing image column cause duh
    available_column = sorted(list(set(pd_data.columns.to_list()) - set([ i for i in image_idx])))
  
    context = {'column':available_column,'id':id}
    
    if(is_uploader(req.user)):
        return render(req,'Excel/view/uploader.html',context=context)
    elif(is_manager(req.user)):
        return render(req,'Excel/view/manager.html',context=context)

def search_query(pd_data:pd.DataFrame,column:str,value,available_column) -> tuple[bool,pd.DataFrame]:
        
    if((column and value) and (column in available_column) and (column in pd_data.columns)):
        sample = pd_data[pd_data[column].astype(str).str.contains(value)]
        
        if(sample.empty):
            return (True,pd_data)
        
        pd_data = sample
    
    return (False,pd_data)
        
def table_query(req: HttpRequest, id:str):
    if(req.method=='GET' and req.META.get('HTTP_HX_REQUEST') and (req.user.is_authenticated)):
        column = req.GET.get('column',None)
        value = req.GET.get('search',None)
        
        result = get_data(id,req.user.username)
        
        pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
        image_idx:list = result.get('excel',{}).get('image',[])
        verify_idx:list = result.get('excel',{}).get('verify',{})
        
        available_column = sorted(list(set(pd_data.columns.to_list()) - set([ i for i in image_idx])))
        
        empty_search, pd_data = search_query(pd_data,column,value,available_column)
        
        context = {'column':pd_data.columns,'result':pd_data.iterrows(),'image':image_idx,'verify':verify_idx}
        
        if(empty_search):
            context.update({'alert':'Data Search returned 0 results'})
        
        return render(req,'HTMX/table.html',context=context)
    
    
def verify_page(req: HttpRequest, id:str, index:int) -> HttpResponse:
    if(not (req.user.is_authenticated)):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    result = get_data(id,req.user.username)
    
    if(result is None):
        return redirect(reverse('Excel:view',kwargs={'id':id}))
    
    pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
    image_idx:list = result.get('excel',{}).get('image',[])
    manager:str = result.get('verify',[])
    verify_idx = result.get('excel',{}).get('verify',{})
    feed_idx = result.get('excel',{}).get('feedback',{})
    lock_idx = result.get('excel',{}).get('locked',{})
    
    if(int(index)>pd_data.shape[0]):
        return redirect(reverse('Excel:view',args={'id':id}) + '?alert=Index not found!')
    
    data = pd_data.iloc[int(index)].to_dict()
    context = {'id':id,'data':data,'image':image_idx,'column':pd_data.columns,**get_post(req.user)}
    
    if(req.method=='POST'):
        f = VerifyForm(req.POST)
        
        if(f.is_valid()):
            status, feed = f.cleaned_data.get("status"), f.cleaned_data.get("feedback") 
            
            if(status=='True'):
                status = True
            elif(status=='False'):
                status = False
            else:
                status = None
                
            verify_idx[str(index)] = status
            feed_idx[str(index)] = feed
            
            excel = MongoConnection(settings.MONGO_URL).connect(settings.MONGO_CRED)    
            excel.update_one({"_id":ObjectId(id)},{"$set":{"excel.verify":verify_idx,'excel.feedback':feed_idx}})
    else:
        f = VerifyForm(initial={'status':verify_idx[str(index)],'feedback':feed_idx[str(index)]})
    
    if(is_uploader(req.user)):
        if(manager is not None):
            context.update({'manager':[{'user':i.get('user',None),'status':verify_idx[str(index)],'feedback':feed_idx[str(index)]} for i in manager],'lock':lock_idx[str(index)]})
            
        return render(req,'Excel/single/uploader.html',context=context)
    
    else:
        context.update({'form':f,'status':verify_idx[str(index)],'lock':lock_idx[str(index)]})
        return render(req,'Excel/single/manager.html',context=context)

def single_query(req: HttpRequest, id:str, index:int):
    if(req.method=='GET' and req.META.get('HTTP_HX_REQUEST') and (req.user.is_authenticated)):
            
        result = get_data(id,req.user.username)
        
        pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
        image_idx:list = result.get('excel',{}).get('image',[])

        if(int(index)>pd_data.shape[0]):
            return redirect(reverse('Excel:view',args={'id':id}) + '?alert=Index not found!')

        else:
            data = pd_data.iloc[int(index)].to_dict()
            
        view = report_view(req.user.username,id,index)
        
        return render(req,'HTMX/report.html',context={'id':id,'data':data,'personal':view.personal_info,'pic':view.profile_img,'sem_dict':view.sem_data,**get_post(req.user)})

def cancel_issue(req: HttpRequest, id:str, index:int):
    if(not (req.user.is_authenticated and is_manager(req.user))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    excel = MongoConnection(settings.MONGO_URL).connect(settings.MONGO_CRED)
    result = get_data(id,req.user.username)
    
    lock_status = result.get('excel',{}).get('locked',{})
    lock_status[str(index)] = {'status':False,'time':None}
    excel.update_one({"_id":ObjectId(id)},{"$set":{'excel.locked':lock_status}})
    
    for locked in IssuedData.objects.filter(mongo_id__mongo_id=id,row_index=index):
        locked.delete()
        
    return redirect(reverse('Excel:single',kwargs={'id':id,'index':index}))

def compress_data(req: HttpRequest, id:str, index:int):
    if(not (req.user.is_authenticated and is_manager(req.user))):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    excel = MongoConnection(settings.MONGO_URL).connect(settings.MONGO_CRED)
    result = get_data(id,req.user.username)
    pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
    image_idx:list = result.get('excel',{}).get('image',[])
    lock_status = result.get('excel',{}).get('locked',{})
    timestamp = timezone.now().isoformat()
    belongs = result.get('belongs',None)
    
    profile_img = r'^Profile_Image$'
    sem_data = r'.+Sem_(\d+)$'
    other_data = """ Anything not part above is personal """
    
    columns = pd_data.keys()
    
    profile_col = [i for i in columns if re.match(profile_img,i)]
    sem_col = [i for i in columns if re.match(sem_data,i)]
    personal_col = [i for i in columns if((i not in profile_col) and (i not in sem_col))]
    
    sem_dict:dict[str,list[str]] = {}
    
    profile_col = profile_col[0] if profile_col else None
    
    for i in sem_col:
        sem:list[str] = re.findall(sem_data,i)
        
        if(sem):
            sem = sem[0]
            if(sem not in sem_dict): sem_dict[sem] = list()
            sem_dict[sem].append(i)


    if(int(index)>pd_data.shape[0]):
        return redirect(reverse('Excel:view',kwargs={'id':id}) + '?alert=Index not found!')

    lock_status[str(index)] = {'status':True,'time':timestamp}
    data = pd_data.iloc[int(index)].to_dict()
    excel.update_one({"_id":ObjectId(id)},{"$set":{'excel.locked':lock_status}})
    
    idx = VerificationTable.objects.filter(Q(mongo_id=id))
    
    if idx:
        idx = idx[0]
    else:
        return redirect(reverse('Excel:dash')+'?alert=Record not found')
    
    locked_data = IssuedData.objects.filter(mongo_id=idx,row_index=index)
    
    if(not locked_data):
        locked = IssuedData(mongo_id=idx,issued=req.user.manager,row_index=index,timestamp=timestamp)
        locked.save()
    else:
        for locked in IssuedData.objects.filter(mongo_id=idx,row_index=index):
            locked.timestamp=timestamp
            locked.save()
    
    for i in image_idx:
        img_data = data[i]
        
        raw_img = img_data.split(':')
    
        try:
            width, height, img = raw_img
        except:
            width, height = 100,100
            img = bad_image
            
        img = compress_image(BytesIO(b64decode(img)))
        
        data[i] = img
        
    send_data = {'header':{'user':req.user.username,**get_post(req.user),'belongs':belongs,'time':timestamp},'data':data}
    print(send_data)
    with open(settings.MEDIA_ROOT + '/compress.txt','w') as f:
        f.write(encrypt_data(settings.KEY,send_data))
        
    return redirect(reverse('Excel:single',kwargs={'id':id,'index':index}))
    
def quick_query(req: HttpRequest, id:str):
    if(req.method=='GET' and req.META.get('HTTP_HX_REQUEST') and (req.user.is_authenticated)):
        column = req.GET.get('column',None)
        search = req.GET.get('search','')
        

        result = get_data(id,req.user.username)

        pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
        
        context = {}

        if((column is None) or (column not in pd_data.columns)):
            return render(req,'HTMX/suggests.html',context=context)
        
        pd_data = pd_data[pd_data[column].str.contains(search)][column].to_numpy()
                
        context.update({'option':[i for i in sorted(set(pd_data))[:5]]})
        
        return render(req,'HTMX/suggests.html',context=context)
        


def get_data(mongo_id:str, username:str) -> dict[str,Any | dict | list]:
    excel = MongoConnection(settings.MONGO_URL).connect(settings.MONGO_CRED)

    result:dict[str,Any | dict | list] = excel.find_one({"$and":[{"_id":ObjectId(mongo_id)},{"$or":[{"belongs":username},{"verify.user":username}]}]})

    
    return result

class ReportStructure(typing.NamedTuple):
    profile_img:Image
    personal_info:dict[str,str]
    sem_data:dict[dict[str,str]]

def report_view(username:str, id:str, index:int) -> ReportStructure:
    
    profile_img = r'^Profile_Image$'
    sem_data = r'.+Sem_(\d+)$'
    other_data = """ Anything not part above is personal """
    
    result:dict[str,Any | dict | list] = get_data(id,username)
    
    pd_data = pd.DataFrame(result.get('excel',{}).get('data',{}))
    
    image_idx:list = result.get('excel',{}).get('image',[])
    single_data = pd_data.iloc[index].to_dict()
    columns = single_data.keys()
    
    profile_col = [i for i in columns if re.match(profile_img,i)]
    sem_col = [i for i in columns if re.match(sem_data,i)]
    personal_col = [i for i in columns if((i not in profile_col) and (i not in sem_col))]
    
    sem_dict:dict[str,list[str]] = {}
    
    profile_col = profile_col[0] if profile_col else None
    
    for i in sem_col:
        sem:list[str] = re.findall(sem_data,i)
        
        if(sem):
            sem = sem[0]
            if(sem not in sem_dict): sem_dict[sem] = list()
            sem_dict[sem].append(i)

    return ReportStructure(profile_img=profile_col,personal_info=personal_col,sem_data=sem_dict)


def edit_view(req: HttpRequest, id:str) -> HttpResponse:
    if(not (req.user.is_authenticated)):
        return redirect(reverse(settings.LOGIN_URL) + '?alert=Unauthenticated Request')
    
    if(not is_uploader(req.user)):
        return redirect(reverse('Excel:view',kwargs={'id':id}) + '?alert=Must be a uploader')
    
    context = {'id':id}
    
    result = get_data(id,req.user.username)
    
    is_locked = any([i.get('status',None) for i in result.get('excel',{}).get('locked',{}).values()])
    
    if(is_locked):
        return redirect(reverse('Excel:view',kwargs={'id':id}) + '?alert=Data is locked')
    
    file_name = result.get('excel',{}).get('name','')
    
    if(req.method=='GET'):
        f = ExcelForm(initial={'username':req.user.username,'file_name':file_name})
        context.update({'form':f})
        
    elif(req.method=='POST'):
        
        f = ExcelForm(req.POST,req.FILES)
        
        context.update({'form':f})
        
        connection = MongoConnection(settings.MONGO_URL)
        excel = connection.connect(settings.MONGO_CRED)
        pd_data = None
        
        if(excel is None):
            return render(req,'Excel/upload.html',{'form':f,'alert':'MongoDB connection failed'})
        
        if(f.is_valid()):
            excel_file = req.FILES["file"]
            file_name = f.cleaned_data.get("file_name")
        
            with excel_file.open() as file:
                image_idx, pd_data = image_load(file.read())
                rows, _ = pd_data.shape
                pd_data = pd_data.to_json()
            
            if(pd_data is None):
                return mongo_setup_failed(req,form=f)
            
            verify_idx:list[dict[str,str]] = result.get('excel',{}).get('verify',[])
            feedback_idx = {str(i):"" for i in range(rows)}
            lock = {str(i):{'status':False,'time':None} for i in range(rows)}
            
            template = MongoTemplate(file_name,json.loads(pd_data),req.user.username,verify_idx,feedback_idx,lock,image_idx).get_json()
            excel.update_one({"_id":ObjectId(id)},{"$set":template})
            
            verify = VerificationTable(mongo_id=id,belongs=req.user.uploader,assigned=(lambda x: x[0].get('user',None) if x else None)(verify_idx))
            verify.save()
            
            context.update({'alert':'Data Updation successful'})
        
        
            
        return render(req,'Excel/edit.html',context=context)
            
    
    return render(req,'Excel/edit.html',context=context)
