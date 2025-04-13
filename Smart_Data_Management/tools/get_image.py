import openpyxl
from PIL import Image
import openpyxl.worksheet
import openpyxl.worksheet.worksheet
from base64 import b64decode, b64encode
from io import BytesIO
import pandas

REDUCE_FACTOR:int = 4

def get_image_data(sheet:openpyxl.worksheet.worksheet.Worksheet) -> dict[tuple[int,int],str]:
    images = {}
    for image in sheet._images:        
        img = Image.open(BytesIO(image._data()))
        
        with BytesIO() as b:
            img.save(b,format=img.format,quality=95)
            images[(image.anchor._from.row,image.anchor._from.col)] = f"{img.width}:{img.height}:{b64encode(b.getvalue()).decode()}"
        
    return images

def compress_image(img_data:BytesIO) -> str:
    image = Image.open(img_data)
    width, height = image.width, image.height
    compressed = image.reduce(REDUCE_FACTOR)
    
    image_data = BytesIO()
    
    compressed.save(image_data,format=image.format,quality=75, optimize=True)
    return f"{width}:{height}:{b64encode(image_data.getvalue()).decode()}"

def expand_image(img_data:str, width:int, height:int) -> str:
    image = Image.open(BytesIO(b64decode(img_data)))
    
    temp = image.resize((width,height),Image.Resampling.BICUBIC)
    image_data = BytesIO()
    temp.save(image_data,format=image.format,quality=95)
    return f"{width}:{height}:{b64encode(image_data.getvalue()).decode()}"

def image_load(data: bytes) -> tuple[list[str], pandas.DataFrame]:
    
    pd_data:pandas.DataFrame = None
    op_data:openpyxl.Workbook = None
    converted:set[str] = set()
    
    try:
        pd_data = pandas.read_excel(BytesIO(data))
        op_data = openpyxl.load_workbook(BytesIO(data))
    except:
        return converted, pd_data
        

    if(op_data and len(op_data.sheetnames)==0):
        return converted, pd_data
        
    op_sheet = op_data[op_data.sheetnames[0]]
    image = get_image_data(op_sheet)
    
    for i, row in pd_data.iterrows():
        for j, _ in enumerate(row):
            
            pd_data[pd_data.columns[j]] = pd_data[pd_data.columns[j]].astype(str)
            if((i+1,j) in image):
                converted.add(pd_data.columns[j])
                    
                pd_data.iat[i,j] = image[(i+1,j)]
                    
    return list(converted), pd_data
