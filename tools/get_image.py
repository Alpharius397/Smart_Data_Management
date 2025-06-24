import openpyxl  # type: ignore
from PIL import Image  # type: ignore
import openpyxl.worksheet  # type: ignore
import openpyxl.worksheet.worksheet  # type: ignore
from tools.encrypt import b64decode, b64encode
from io import BytesIO
import pandas
from tools.errors import (
    ImageCompressionFailed,
    ImageExpansionFailed,
)

REDUCE_FACTOR: int = 5


def get_image_data(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
) -> dict[tuple[int, int], str]:
    images = {}
    for image in sheet._images:  # type: ignore
        img = Image.open(BytesIO(image._data()))
        width, height = img.width, img.height
        default_size = 200
        
        if max(width, height) > 200:
            if width > height:
                height = int((height / width) * default_size) 
                width = default_size
            else:
                width = int((width / height) * default_size)
                height = default_size

        temp = img.resize((width, height), Image.Resampling.BILINEAR)

        with BytesIO() as b:
            temp.save(b, format=img.format, quality=95)
            images[(image.anchor._from.row, image.anchor._from.col)] = b64encode(b.getvalue()).decode()
            

    return images


def compress_image(img_data: str) -> str:
    try:
        image = Image.open(BytesIO(b64decode(img_data)))
        compressed = image.reduce(REDUCE_FACTOR)
        image_data = BytesIO()
        compressed.save(image_data, format=image.format, quality=75, optimize=True)
        
        with open("/home/omnissiah/Project/nodejs/react/Smart_Data_Management/sample/compressed_image.txt", "w") as f:
            f.write(b64encode(image_data.getvalue()).decode())
            
        with open("/home/omnissiah/Project/nodejs/react/Smart_Data_Management/sample/image.txt", "w") as f:
            f.write(img_data)
            
        return b64encode(image_data.getvalue()).decode()

    except Exception:
        raise ImageCompressionFailed()


def expand_image(img_data: str, width: int, height: int) -> str:
    try:
        image = Image.open(BytesIO(b64decode(img_data)))

        temp = image.resize((width, height), Image.Resampling.BILINEAR)
        image_data = BytesIO()
        temp.save(image_data, format=image.format, quality=95)
        return b64encode(image_data.getvalue()).decode()
    except Exception:
        raise ImageExpansionFailed()


def image_load(data: bytes) -> tuple[list[int], pandas.DataFrame]:
    pd_data: pandas.DataFrame = pandas.DataFrame()
    op_data: openpyxl.Workbook = openpyxl.Workbook()
    converted: set[int] = set()

    try:
        op_data = openpyxl.load_workbook(BytesIO(data))
        pd_data = pandas.read_excel(BytesIO(data))
    except Exception as e:
        return [], pd_data

    if op_data and len(op_data.sheetnames) == 0:
        return [], pd_data

    op_sheet = op_data[op_data.sheetnames[0]]
    image = get_image_data(op_sheet)

    rows, columns = pd_data.shape

    for i in range(rows):
        for j in range(columns):
            if (i + 1, j) in image:
                converted.add(j)

                pd_data.iat[i, j] = image[(i + 1, j)]

    return list(converted), pd_data
